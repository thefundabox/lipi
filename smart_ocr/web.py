"""Local web UI: upload a PDF, pick input type and language, watch progress, get the text.

Run:  .venv/bin/python -m smart_ocr.web      then open http://127.0.0.1:8000
"""
import argparse
import io
import os
import re
import tempfile
import threading
import uuid
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pymupdf
from docx import Document
from docx.shared import Pt
from flask import Flask, Response, abort, jsonify, request, send_file, send_from_directory

from . import engines
from .cli import parse_pages
from .converter import convert
from .profiles import LANGUAGES, PROFILES

STATIC = Path(__file__).parent / "web_static"
MAX_UPLOAD_MB = 200
KEEP_JOBS = 10  # uploaded PDFs are kept (for page previews) for the most recent jobs only

app = Flask(__name__, static_folder=None)
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_MB * 1024 * 1024

# OCR is CPU-heavy, so jobs run one at a time; later uploads wait in the queue.
executor = ThreadPoolExecutor(max_workers=1)
jobs: dict[str, dict] = {}
jobs_lock = threading.Lock()


def _update(job_id, **fields):
    with jobs_lock:
        jobs[job_id].update(fields)


def _run(job_id, path, profile, lang, allow_cloud, pages, fix_spelling):
    def on_page(result, index, total):
        with jobs_lock:
            job = jobs[job_id]
            job["pages"].append(result.__dict__.copy())
            job["done"], job["total"] = index, total

    try:
        _update(job_id, status="running")
        convert(path, profile, lang, allow_cloud, pages, log=lambda *_: None, on_page=on_page,
                fix_spelling=fix_spelling)
        _update(job_id, status="done")
    except Exception as e:  # surface any failure to the page instead of a silent hang
        _update(job_id, status="error", error=f"{type(e).__name__}: {e}")


def _evict_old_jobs():
    with jobs_lock:
        finished = [j for j in jobs.values() if j["status"] in ("done", "error")]
        for job in finished[:max(0, len(finished) - KEEP_JOBS)]:
            try:
                Path(job["_path"]).unlink(missing_ok=True)
            except OSError:  # Windows: still open by a preview request; the temp folder is cleaned later
                pass
            del jobs[job["id"]]


def _get_job(job_id):
    with jobs_lock:
        job = jobs.get(job_id)
        if job is None:
            abort(404)
        return {k: v for k, v in job.items() if not k.startswith("_")}, job["_path"]


@app.get("/")
def index():
    return send_from_directory(STATIC, "index.html")


@app.get("/static/<path:name>")
def static_file(name):
    return send_from_directory(STATIC, name)


@app.get("/api/options")
def options():
    status = engines.availability()
    return jsonify({
        "types": [{"key": p.key, "label": p.label, "description": p.description,
                   "needs_cloud": p.engines == ("claude",)} for p in PROFILES.values()],
        "languages": [{"key": k, "label": v[0]} for k, v in LANGUAGES.items()],
        "cloud_available": status["claude"] == "available",
        "cloud_status": status["claude"],
        "max_upload_mb": MAX_UPLOAD_MB,
    })


@app.post("/api/jobs")
def create_job():
    upload = request.files.get("file")
    profile = PROFILES.get(request.form.get("type", ""))
    lang = request.form.get("lang", "")
    if not upload or not upload.filename:
        return jsonify(error="Choose a PDF file."), 400
    if profile is None or lang not in LANGUAGES:
        return jsonify(error="Unknown input type or language."), 400
    try:
        pages = parse_pages(request.form["pages"]) if request.form.get("pages", "").strip() else None
    except ValueError:
        return jsonify(error="Pages should look like 1-3,7"), 400

    fd, path = tempfile.mkstemp(suffix=".pdf")
    with os.fdopen(fd, "wb") as f:
        upload.save(f)
    try:
        with pymupdf.open(path) as doc:
            if not doc.is_pdf:
                raise ValueError
            page_count = doc.page_count
    except Exception:
        os.unlink(path)
        return jsonify(error="That file isn't a readable PDF."), 400

    total = len([n for n in pages if 1 <= n <= page_count]) if pages else page_count
    if total == 0:
        os.unlink(path)
        return jsonify(error=f"No such pages; this PDF has {page_count}."), 400

    job_id = uuid.uuid4().hex
    with jobs_lock:
        jobs[job_id] = {"id": job_id, "filename": upload.filename, "type": profile.label,
                        "language": LANGUAGES[lang][0], "lang": lang, "status": "queued",
                        "done": 0, "total": total, "pages": [], "error": None, "_path": path}
    allow_cloud = request.form.get("allow_cloud") == "1" and engines.availability()["claude"] == "available"
    _evict_old_jobs()
    fix_spelling = request.form.get("spellfix", "1") == "1"
    executor.submit(_run, job_id, path, profile, lang, allow_cloud, pages, fix_spelling)
    return jsonify(id=job_id)


@app.get("/api/jobs/<job_id>")
def job_status(job_id):
    return jsonify(_get_job(job_id)[0])


@app.get("/api/jobs/<job_id>/page/<int:number>.png")
def page_image(job_id, number):
    _, path = _get_job(job_id)
    with pymupdf.open(path) as doc:
        if not 1 <= number <= doc.page_count:
            abort(404)
        png = doc[number - 1].get_pixmap(dpi=110).tobytes("png")
    return Response(png, mimetype="image/png", headers={"Cache-Control": "max-age=3600"})


@app.post("/api/export/docx")
def export_docx():
    """Build a Word file from the (possibly edited) page texts the browser sends back."""
    data = request.get_json(silent=True) or {}
    pages = data.get("pages") or []
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Nirmala UI", Pt(12)  # renders Devanagari and Latin in Word
    for i, page in enumerate(pages):
        if i:
            doc.add_page_break()
        doc.add_heading(f"Page {page.get('number', i + 1)}", level=2)
        for para in str(page.get("text", "")).split("\n"):
            doc.add_paragraph(para)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    stem = re.sub(r"[^\w\-. ]", "_", Path(str(data.get("filename", "converted"))).stem) or "converted"
    return send_file(buf, as_attachment=True, download_name=f"{stem}.docx",
                     mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


@app.post("/api/quit")
def quit_app():
    """Stop Lipi. Used by the Quit button, since the packaged app has no window of its own."""
    threading.Timer(0.3, lambda: os._exit(0)).start()
    return jsonify(ok=True)


@app.errorhandler(413)
def too_large(_):
    return jsonify(error=f"File is larger than {MAX_UPLOAD_MB} MB."), 413


def main():
    ap = argparse.ArgumentParser(description="Web UI for smart_ocr")
    ap.add_argument("--port", type=int, default=8000)
    args = ap.parse_args()
    # Bound to this computer only: uploads and results never leave the machine unless cloud is allowed.
    print(f"Open http://127.0.0.1:{args.port} in your browser (Ctrl+C to stop)")
    app.run(host="127.0.0.1", port=args.port, threaded=True)


if __name__ == "__main__":
    main()
